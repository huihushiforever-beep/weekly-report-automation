#!/usr/bin/env python3
"""
Re-opens the spreadsheet that was just written and checks the numbers in it
against a fresh calculation from the raw exports. Run it after build_report.py.

    python verify_output.py

It exits with code 1 and prints DIFF if anything does not match.
"""

from __future__ import annotations

import sys
from datetime import timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def recompute(orders_path: str, accounts_path: str) -> dict:
    orders = pd.read_csv(orders_path, dtype={"account_id": str, "qty": str})
    accounts = pd.read_csv(accounts_path, dtype={"account_id": str})

    def parse(value):
        for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
            try:
                return pd.to_datetime(str(value).strip(), format=pattern)
            except (ValueError, TypeError):
                continue
        return pd.NaT

    orders["date"] = orders["order_date"].apply(parse)
    orders["qty"] = pd.to_numeric(orders["qty"], errors="coerce")
    orders["unit_price"] = pd.to_numeric(orders["unit_price"], errors="coerce")
    orders = orders.dropna(subset=["date", "qty"])
    orders["revenue"] = orders["qty"] * orders["unit_price"].fillna(0.0)
    orders = orders.merge(accounts[["account_id"]], on="account_id", how="inner")

    last_day = orders["date"].max()
    start = last_day - timedelta(days=6)
    week = orders[(orders["date"] >= start) & (orders["date"] <= last_day)]
    return {
        "revenue": round(float(week["revenue"].sum()), 2),
        "orders": int(len(week)),
        "units": int(week["qty"].sum()),
        "accounts": int(week["account_id"].nunique()),
    }


def main() -> int:
    out_dir = Path("output")
    reports = sorted(out_dir.glob("weekly_report_*.xlsx"))
    if not reports:
        print("no report found in output/ - run build_report.py first")
        return 1
    path = reports[-1]
    expected = recompute("sample_data/sales_orders.csv", "sample_data/crm_accounts.csv")

    sheet = load_workbook(path, data_only=True)["Summary"]
    found = {}
    for row in range(5, 10):
        label = sheet.cell(row=row, column=1).value
        if label:
            found[label] = sheet.cell(row=row, column=2).value

    checks = [
        ("Revenue", round(float(found.get("Revenue", 0)), 2), expected["revenue"]),
        ("Orders", int(found.get("Orders", 0)), expected["orders"]),
        ("Units sold", int(found.get("Units sold", 0)), expected["units"]),
        ("Active accounts", int(found.get("Active accounts", 0)), expected["accounts"]),
    ]

    problems = 0
    print(f"checking {path}")
    for label, in_file, recomputed in checks:
        status = "OK  " if in_file == recomputed else "DIFF"
        if in_file != recomputed:
            problems += 1
        print(f"  {status} {label:<18} file={in_file:<14} recomputed={recomputed}")

    product_sheet = load_workbook(path, data_only=True)["By product"]
    product_total = round(sum(product_sheet.cell(row=r, column=4).value or 0
                              for r in range(4, product_sheet.max_row + 1)), 2)
    if product_total != expected["revenue"]:
        problems += 1
        print(f"  DIFF product sheet total={product_total} recomputed={expected['revenue']}")
    else:
        print(f"  OK   product sheet totals to {product_total}")

    print("result:", "all checks passed" if problems == 0 else f"{problems} mismatch(es)")
    return 0 if problems == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
